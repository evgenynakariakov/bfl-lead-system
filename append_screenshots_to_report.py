from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


def add_paragraph_justified(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.line_spacing = 1.5


def add_heading(doc: Document, text: str, size_pt: int = 13) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size_pt)


def add_figure(doc: Document, image_path: Path, caption: str) -> None:
    pic_p = doc.add_paragraph()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.add_run().add_picture(str(image_path), width=Cm(16.2))

    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(11)


def main() -> None:
    root = Path(r"D:\bfl_system_complete")
    report_dir = root / "отчет"
    report_path = report_dir / "Ответ на тестовое задание.docx"
    backup_path = report_dir / "Ответ на тестовое задание_backup_before_screens.docx"

    if not report_path.exists():
        raise FileNotFoundError(f"Не найден файл отчета: {report_path}")

    if not backup_path.exists():
        original = Document(str(report_path))
        original.save(str(backup_path))

    doc = Document(str(report_path))

    # Guard: do not append duplicate block if already exists.
    if any(
        (p.text or "").strip().startswith(
            "8. Подтверждение работоспособности системы на тестовом прогоне"
        )
        for p in doc.paragraphs
    ):
        print("Блок со скриншотами уже присутствует, повторно не добавляю.")
        return

    doc.add_page_break()
    add_heading(doc, "8. Подтверждение работоспособности системы на тестовом прогоне")
    add_paragraph_justified(
        doc,
        "Ниже приведены скриншоты реального прогона по этапам воронки: запуск парсинга, "
        "генерация первого касания, формирование персональных deep-link, прохождение "
        "квалификационного сценария в Telegram-боте и передача квалифицированного лида менеджеру. "
        "Блок добавлен как доказательная часть отчета: каждый шаг, который был описан в архитектуре "
        "и UML, подтвержден фактическими логами и данными в базе.",
    )

    figures = [
        (
            report_dir / "Запуск парсера.png",
            "Рисунок 4. Запуск парсера Авито и логирование статуса выполнения.",
        ),
        (
            report_dir / "запуск блока первого касания.png",
            "Рисунок 5. Массовая обработка лидов в сервисе первого касания (sent=29, error=0).",
        ),
        (
            report_dir / "Сгенерированный текст первого касание и уникальная ссылка телеграм.png",
            "Рисунок 6. Сохраненные в БД first_touch_text, outreach_status и first_touch_at для каждого лида.",
        ),
        (
            report_dir / "персональная ссылка (Deeplink).png",
            "Рисунок 7. Персональные deep-link для перехода лида в Telegram-бота.",
        ),
        (
            report_dir / "запуск бота и логи.png",
            "Рисунок 8. Логи Telegram-бота с переходом между шагами FSM и уведомлением менеджеру.",
        ),
        (
            report_dir / "тест в тгботе.png",
            "Рисунок 9. Диалог квалификации в Telegram-боте по сценарию debt -> income -> property -> creditors.",
        ),
        (
            report_dir / "результат работы бота, информация которая преедаётся Оператору.png",
            "Рисунок 10. Карточка квалифицированного лида, переданная менеджеру.",
        ),
        (
            report_dir / "скоринг.png",
            "Рисунок 11. Таблица лидов в БД с полями score и score_breakdown после этапа скоринга.",
        ),
    ]

    for image_path, caption in figures:
        if image_path.exists():
            add_figure(doc, image_path, caption)
        else:
            miss = doc.add_paragraph(f"{caption} (файл {image_path.name} не найден)")
            miss.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(doc, "9. Результат сквозного прогона")
    add_paragraph_justified(
        doc,
        "Скриншоты подтверждают, что система работает как единый конвейер. Сначала формируется "
        "входящий поток объявлений и выполняется скоринг, затем создаются персональные тексты и deep-link, "
        "после чего лид проходит автоматическую квалификацию в боте. На выходе менеджер получает "
        "структурированную карточку с параметрами лида, а в базе данных остается полный цифровой след "
        "процесса, пригодный для аналитики, расчета CTR/CPL/ROI и последующей оптимизации сценариев.",
    )

    doc.save(str(report_path))
    print(f"Готово: {report_path}")
    print(f"Backup: {backup_path}")


if __name__ == "__main__":
    main()
